"""Extract released Figure 4 electrical traces, without inventing intake labels."""
import csv
import hashlib
import json
from pathlib import Path
from statistics import median

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / 'dataset/polycore_lipid_2026'
MEALS = {'4e': 'fasting', '4f': 'bacon_and_eggs', '4g': 'pizza',
         '4h': 'ground_beef', '4i': 'milkshake'}


def main():
    source = SOURCE / '44460_2026_117_MOESM5_ESM.xlsx'
    wb = load_workbook(source, read_only=True, data_only=True)
    summaries = []
    for sheet, meal in MEALS.items():
        rows = list(wb[sheet].values)
        header = rows[2]
        start = header.index('Time (s)')
        names = ['sensor_time_s', 'glucose_current_nA', 'cholesterol_current_nA',
                 'tg_proxy_current_nA', 'ph_voltage_as_recorded']
        values = [r[start:start+5] for r in rows[3:]
                  if isinstance(r[start], (int, float))]
        assert values and all(len(r) == 5 for r in values)
        assert all(values[i][0] > values[i-1][0] for i in range(1, len(values)))
        with (OUT / f'{sheet}_{meal}_electrical.csv').open('w') as f:
            writer = csv.writer(f)
            writer.writerow(['source_excel_row', *names])
            for i, row in enumerate(rows[3:], 4):
                if isinstance(row[start], (int, float)):
                    writer.writerow([i, *row[start:start+5]])
        stats = {'sheet': sheet, 'condition': meal, 'rows': len(values),
                 'first_time_s': values[0][0], 'last_time_s': values[-1][0],
                 'ph_voltage_header': header[start+4],
                 'known_consumed_macro_grams': None,
                 'participant_id': None,
                 'note': 'Post-placement electrical signal, not premeal baseline or intake grams.'}
        for i, name in enumerate(names[1:4], 1):
            col = [r[i] for r in values]
            assert all(isinstance(v, (int, float)) for v in col)
            first = median([r[i] for r in values if r[0] <= values[0][0]+59])
            last = median([r[i] for r in values if r[0] >= values[-1][0]-59])
            stats[name] = {'first_minute_median': first, 'last_minute_median': last,
                           'minimum': min(col), 'maximum': max(col)}
        summaries.append(stats)
    wb.close()
    metric_wb = load_workbook(SOURCE / '44460_2026_117_MOESM6_ESM.xlsx',
                             read_only=True, data_only=True)
    metrics = [list(r) for r in metric_wb['5hi'].values]
    metric_wb.close()
    result = {'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
              'meal_traces': summaries, 'figure5hi_workbook_rows': metrics,
              'interpretation': 'Five plotted condition traces are not five independent dose cohorts. '
              'No participant grouping or consumed-dose labels inferred from figure order.'}
    (OUT / 'extraction_summary.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
